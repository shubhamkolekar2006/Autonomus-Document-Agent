import io
import os
import re
from typing import BinaryIO, Dict, List, Optional, Union
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

# Configurable max file size limit (default 10 MB)
MAX_FILE_SIZE_MB = float(os.getenv("MAX_FILE_SIZE_MB", "10.0"))
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class DocumentParsingError(Exception):
    """Base exception for document parsing errors."""
    pass


class UnsupportedFileTypeError(DocumentParsingError):
    """Raised when an uploaded file extension is not supported."""
    pass


class FileTooLargeError(DocumentParsingError):
    """Raised when an uploaded file exceeds the allowed size limit."""
    pass


class EmptyDocumentError(DocumentParsingError):
    """Raised when an uploaded file is empty or contains no extractable text."""
    pass


class CorruptedDocumentError(DocumentParsingError):
    """Raised when an uploaded file is corrupted or cannot be read."""
    pass


def normalize_text(text: str) -> str:
    """
    Normalizes extracted text by:
    - Normalizing newlines (\r\n and \r to \n)
    - Removing non-printable null bytes
    - Replacing excessive horizontal whitespace with a single space
    - Collapsing 3+ consecutive newlines to 2 newlines (preserving paragraph structure)
    - Stripping leading and trailing whitespace
    """
    if not text:
        return ""
    
    # Replace null bytes
    text = text.replace("\x00", "")
    
    # Normalize unicode hyphens, dashes, and curly quotes
    unicode_replacements = {
        "\u2011": "-",  # non-breaking hyphen
        "\u2012": "-",  # figure dash
        "\u2013": "-",  # en dash
        "\u2014": "--", # em dash
        "\u2018": "'",  # left single quote
        "\u2019": "'",  # right single quote
        "\u201c": '"',  # left double quote
        "\u201d": '"',  # right double quote
        "\u2026": "...", # ellipsis
        "\u00a0": " ",  # non-breaking space
    }
    for orig, rep in unicode_replacements.items():
        text = text.replace(orig, rep)

    # Standardize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    
    # Normalize horizontal whitespace on each line (tabs/spaces) without destroying newlines
    lines = []
    for line in text.split("\n"):
        cleaned_line = re.sub(r"[^\S\n]+", " ", line).strip()
        lines.append(cleaned_line)
    
    rebuilt = "\n".join(lines)
    
    # Collapse 3 or more consecutive newlines into 2 (blank line between paragraphs)
    rebuilt = re.sub(r"\n{3,}", "\n\n", rebuilt)
    
    return rebuilt.strip()


def parse_text(content: Union[bytes, str]) -> str:
    """
    Parses plain text (.txt) content safely using UTF-8 with fallback decoding.
    """
    if isinstance(content, bytes):
        if not content:
            raise EmptyDocumentError("Text file is empty.")
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception as e:
                raise CorruptedDocumentError(f"Failed to decode text file: {str(e)}") from e
    else:
        text = content
    
    normalized = normalize_text(text)
    if not normalized:
        raise EmptyDocumentError("Text file contains no readable text.")
    return normalized


def parse_markdown(content: Union[bytes, str]) -> str:
    """
    Parses Markdown (.md) content.
    """
    return parse_text(content)


def parse_pdf(content: Union[bytes, BinaryIO]) -> str:
    """
    Parses PDF (.pdf) binary content using pypdf.
    Extracts text page-by-page.
    """
    if isinstance(content, bytes):
        if not content:
            raise EmptyDocumentError("PDF file is empty.")
        stream = io.BytesIO(content)
    else:
        stream = content

    try:
        reader = PdfReader(stream)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:
                raise CorruptedDocumentError(f"PDF is password protected or encrypted: {str(e)}") from e

        pages_text = []
        for i, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    pages_text.append(page_text)
            except Exception as page_err:
                print(f"Warning: Failed to extract text from page {i + 1}: {str(page_err)}")

        combined = "\n\n".join(pages_text)
        normalized = normalize_text(combined)

        if not normalized:
            raise EmptyDocumentError("PDF contains no readable text. It may be empty or a scanned image.")

        return normalized

    except (PdfReadError, Exception) as e:
        if isinstance(e, EmptyDocumentError):
            raise
        raise CorruptedDocumentError(f"Failed to parse PDF document: {str(e)}") from e


def parse_docx(content: Union[bytes, BinaryIO]) -> str:
    """
    Parses Microsoft Word (.docx) binary content using python-docx.
    Extracts text from paragraphs and table cells.
    """
    if isinstance(content, bytes):
        if not content:
            raise EmptyDocumentError("DOCX file is empty.")
        stream = io.BytesIO(content)
    else:
        stream = content

    try:
        doc = Document(stream)
        extracted_elements = []

        # Extract paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                extracted_elements.append(p.text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    extracted_elements.append(" | ".join(row_cells))

        combined = "\n\n".join(extracted_elements)
        normalized = normalize_text(combined)

        if not normalized:
            raise EmptyDocumentError("DOCX contains no readable text.")

        return normalized

    except Exception as e:
        if isinstance(e, EmptyDocumentError):
            raise
        raise CorruptedDocumentError(f"Failed to parse DOCX document: {str(e)}") from e


def parse_document(file_name: str, file_bytes: bytes, max_file_size_mb: float = MAX_FILE_SIZE_MB) -> Dict[str, str]:
    """
    Validates and parses an uploaded document file.
    
    Checks:
    - File size <= max_file_size_mb
    - Supported extensions (.txt, .md, .pdf, .docx)
    - Non-empty payload and non-empty extracted text
    
    Returns structured dictionary:
    {
        "source_name": file_name,
        "source_type": ext (e.g. "pdf", "docx", "txt", "md"),
        "content": extracted_text,
        "character_count": int
    }
    """
    if not file_bytes:
        raise EmptyDocumentError(f"File '{file_name}' is empty (0 bytes).")

    # File size validation
    file_size_mb = len(file_bytes) / (1024 * 1024)
    if file_size_mb > max_file_size_mb:
        raise FileTooLargeError(
            f"File '{file_name}' ({file_size_mb:.2f} MB) exceeds the maximum allowed size limit of {max_file_size_mb} MB."
        )

    # Extension validation
    _, ext = os.path.splitext(file_name.lower())
    if ext not in SUPPORTED_EXTENSIONS:
        supported_str = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileTypeError(
            f"File '{file_name}' has unsupported format '{ext}'. Supported formats: {supported_str}"
        )

    # Dispatch to parser
    clean_ext = ext.lstrip(".")
    if ext in {".txt"}:
        text = parse_text(file_bytes)
    elif ext in {".md"}:
        text = parse_markdown(file_bytes)
    elif ext == ".pdf":
        text = parse_pdf(file_bytes)
    elif ext == ".docx":
        text = parse_docx(file_bytes)
    else:
        raise UnsupportedFileTypeError(f"Unsupported format: {ext}")

    return {
        "source_name": file_name,
        "source_type": clean_ext,
        "content": text,
        "character_count": len(text)
    }


def build_source_context(
    sources: Optional[List[Dict[str, str]]] = None,
    direct_text: Optional[str] = None
) -> Dict:
    """
    Combines parsed document sources and direct text into a structured source context dictionary.
    
    Structure:
    {
        "sources": [
            {"source_name": "...", "source_type": "...", "content": "..."},
            ...
        ],
        "combined_text": "SOURCE: ...\n\n..."
    }
    """
    all_sources = []
    
    if sources:
        for s in sources:
            if s.get("content"):
                all_sources.append({
                    "source_name": s.get("source_name", "uploaded_file"),
                    "source_type": s.get("source_type", "document"),
                    "content": s.get("content", "")
                })

    if direct_text and direct_text.strip():
        normalized_direct = normalize_text(direct_text)
        if normalized_direct:
            all_sources.append({
                "source_name": "user_provided_text",
                "source_type": "text",
                "content": normalized_direct
            })

    formatted_text = format_source_context_for_prompt(all_sources)

    return {
        "sources": all_sources,
        "combined_text": formatted_text
    }


def format_source_context_for_prompt(sources: List[Dict[str, str]]) -> str:
    """
    Formats the list of sources into a standardized string representation for LLM prompts:
    
    SOURCE: <source_name>
    <content>
    """
    if not sources:
        return ""

    blocks = []
    for src in sources:
        name = src.get("source_name", "source")
        content = src.get("content", "").strip()
        if content:
            blocks.append(f"SOURCE: {name}\n{content}")

    return "\n\n" + ("\n\n" + "=" * 40 + "\n\n").join(blocks) + "\n\n"
