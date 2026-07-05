import os
import re
import time
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_page_number(run):
    """
    Appends a dynamic PAGE field XML elements into the given run to render page numbers in Word.
    """
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def add_bottom_border(paragraph):
    """
    Adds a professional bottom accent border (blue) to a paragraph in MS Word.
    """
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')  # Thickness
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '2563EB')  # Accent Blue color
    pBdr.append(bottom)
    pPr.append(pBdr)

def parse_and_add_table(doc, table_lines):
    """
    Parses collected Markdown table lines and adds them as a styled Word table.
    """
    rows_data = []
    for line in table_lines:
        cells = [c.strip() for c in line.split('|')]
        # Strip outer empty cells caused by leading/trailing pipes
        if cells and cells[0] == '':
            cells = cells[1:]
        if cells and cells[-1] == '':
            cells = cells[:-1]
            
        # Skip Markdown header separator line (e.g. |---|---|)
        is_separator = all(re.match(r'^[:\-\s]+$', c) for c in cells) if cells else False
        if not is_separator and len(cells) > 0:
            rows_data.append(cells)
            
    if not rows_data:
        return
        
    num_cols = max(len(row) for row in rows_data)
    num_rows = len(rows_data)
    
    # Create the Word Table
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = 'Light Shading Accent 1'  # Clean, professional built-in style
    
    # Format and populate cells
    for r_idx, row_cells in enumerate(rows_data):
        row = table.rows[r_idx]
        for c_idx, cell_value in enumerate(row_cells):
            if c_idx < len(row.cells):
                cell = row.cells[c_idx]
                cell.text = cell_value
                
                # Format headers to bold
                if r_idx == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.name = 'Calibri'
                            run.font.size = Pt(10.5)

def generate_docx(document_data: dict, output_dir: str = "outputs") -> str:
    """
    Generates a polished Word document optimized for page space efficiency:
    - Creates a cover page with moderate spacing.
    - Adds headers and page number footers on subsequent pages.
    - Automatically maps sections to Heading 1 or Heading 2.
    - All remaining sections flow continuously on subsequent pages with clean spacing.
    - Identifies Markdown tables in content and renders them as styled Word tables.
    - Formats paragraphs and lists cleanly using Calibri.
    """
    os.makedirs(output_dir, exist_ok=True)
    
    doc = Document()
    
    # 1. Page Margins (1 inch margins on all sides)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Configure different first page to suppress header/footer on title page
        section.different_first_page_header_footer = True
        
        # Add page numbering to subsequent pages footer (Right aligned)
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = 2  # Right alignment
        footer_p.text = f"{document_data.get('title', 'Document')} | Page "
        add_page_number(footer_p.add_run())

        # Add document title to subsequent pages header (Right aligned, gray italic)
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = 2  # Right alignment
        header_run = header_p.add_run(document_data.get("title", "Document"))
        header_run.font.name = 'Calibri'
        header_run.font.size = Pt(8.5)
        header_run.font.italic = True
        header_run.font.color.rgb = RGBColor(128, 128, 128)

    # 2. Font Settings (Set Calibri as default for Normal and Headings)
    for style_name in ['Normal', 'Heading 1', 'Heading 2']:
        if style_name in doc.styles:
            doc.styles[style_name].font.name = 'Calibri'

    # 3. Create Cover Page Content (Reduced spacing)
    for _ in range(2):
        doc.add_paragraph()
        
    title_p = doc.add_paragraph()
    title_p.alignment = 1  # Center alignment
    title_run = title_p.add_run(document_data.get("title", "Untitled Document"))
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    
    doc.add_paragraph()  # minor spacer
    
    subtype_p = doc.add_paragraph()
    subtype_p.alignment = 1  # Center alignment
    subtype_run = subtype_p.add_run(document_data.get("document_type", "Document").upper())
    subtype_run.font.size = Pt(13)
    subtype_run.font.italic = True
    add_bottom_border(subtype_p)  # Horizontal separator line
    
    for _ in range(4):
        doc.add_paragraph()
        
    date_p = doc.add_paragraph()
    date_p.alignment = 1  # Center alignment
    date_run = date_p.add_run(f"Generated on {datetime.now().strftime('%B %d, %Y')}")
    date_run.font.size = Pt(10.5)
    
    doc.add_page_break()

    # 4. Add Sections (Flowing continuously on subsequent pages with proper spacing)
    sections = document_data.get("sections", [])
    
    for sec in sections:
        heading = sec.get("heading", "")
        content = sec.get("content", "")
        
        if heading:
            # Check if this heading represents a sub-heading (contains decimal numbering like X.Y)
            is_subheading = re.match(r'^\d+\.\d+', heading.strip()) is not None
            level = 2 if is_subheading else 1
            
            h = doc.add_heading(heading, level=level)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            
        if content:
            lines = content.split('\n')
            in_table = False
            table_lines = []
            
            for line in lines:
                line_str = line.strip()
                if not line_str:
                    if in_table:
                        # End of table block reached
                        parse_and_add_table(doc, table_lines)
                        table_lines = []
                        in_table = False
                    continue
                
                # Check if this line is part of a Markdown table (starts and/or contains pipes)
                is_table_line = line_str.startswith('|') or (line_str.count('|') >= 2)
                
                if is_table_line:
                    in_table = True
                    table_lines.append(line_str)
                else:
                    if in_table:
                        # End of table block reached
                        parse_and_add_table(doc, table_lines)
                        table_lines = []
                        in_table = False
                    
                    # Format as bullet list or normal paragraph
                    is_bullet = line_str.startswith(('-', '*', '•'))
                    
                    if is_bullet:
                        cleaned_line = re.sub(r'^[-*•]\s*', '', line_str)
                        p = doc.add_paragraph(cleaned_line, style='List Bullet')
                        p.paragraph_format.space_after = Pt(3)
                    else:
                        p = doc.add_paragraph(line_str)
                        p.paragraph_format.space_after = Pt(6)
                        p.paragraph_format.line_spacing = 1.15
            
            # If the content ended while in a table block
            if in_table and table_lines:
                parse_and_add_table(doc, table_lines)

    # 5. Save document with slugified title and timestamp
    title_text = document_data.get("title", "document")
    slug = re.sub(r'[^a-zA-Z0-9_-]', '_', title_text.lower())
    slug = re.sub(r'_+', '_', slug).strip('_')
    if not slug:
        slug = "document"
        
    timestamp = int(time.time())
    filename = f"{slug}-{timestamp}.docx"
    file_path = os.path.join(output_dir, filename)
    doc.save(file_path)
    
    return filename

class DocGenerator:
    """
    DocGenerator: Outputs content into formatted MS Word (.docx) files.
    """
    def __init__(self, output_dir: str = "outputs"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    def generate_docx(self, title: str, sections: list, filename: str = "document.docx") -> str:
        """
        Backward-compatible method to generate docx files in self.output_dir.
        """
        document_data = {
            "title": title,
            "document_type": "document",
            "sections": sections
        }
        fname = generate_docx(document_data, output_dir=self.output_dir)
        return os.path.join(self.output_dir, fname)

pre_configured_doc_generator = DocGenerator()
