import io
import json
import unittest
from unittest.mock import patch
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter

from main import app


def mock_call_llm_handler(system_prompt: str, user_prompt: str, json_mode: bool = False) -> str:
    """Mock LLM response for deterministic testing of the ingestion layer."""
    if json_mode:
        if "planning" in system_prompt.lower() or "planner" in system_prompt.lower():
            return json.dumps({
                "document_type": "project plan",
                "assumptions": ["Timeline is 3 months.", "Budget is $75,000."],
                "title": "Grounded Document Plan",
                "sections": [
                    {"heading": "1. Project Overview", "purpose": "Overview grounded in source context."},
                    {"heading": "2. Technical Specifications", "purpose": "Specifications from source context."},
                    {"heading": "3. Implementation & Timeline", "purpose": "Milestones and roadmap."}
                ]
            })
        else:
            # Reflector review call
            return json.dumps({
                "issues_found": [],
                "needs_revision": False
            })
    else:
        # Executor section drafting
        return f"This section is drafted and grounded based on the user prompt and provided sources: {user_prompt[:80]}..."


class TestAgentEndpointIngestion(unittest.TestCase):

    def setUp(self):
        self.client = TestClient(app)

        # Build in-memory sample DOCX
        doc = Document()
        doc.add_paragraph("TechCorp Mobile Banking Roadmap.")
        doc.add_paragraph("Security: OAuth 2.0 with Biometrics authentication.")
        doc.add_paragraph("Launch timeline: Q4 2026.")
        docx_stream = io.BytesIO()
        doc.save(docx_stream)
        self.sample_docx_bytes = docx_stream.getvalue()

    @patch("agent.planner.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.executor.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.reflector.call_llm", side_effect=mock_call_llm_handler)
    def test_scenario_a_no_source_material(self, mock_reflector, mock_executor, mock_planner):
        """Test A: Standard request with JSON payload and no source material."""
        payload = {
            "request": "Create a project plan for launching a mobile banking app in 3 months."
        }
        response = self.client.post("/agent", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("download_url", data)
        self.assertIn("plan", data)
        self.assertIn("times", data)
        self.assertIn("agent_tasks", data)
        self.assertEqual(len(data.get("sources", [])), 0)

    @patch("agent.planner.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.executor.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.reflector.call_llm", side_effect=mock_call_llm_handler)
    def test_scenario_b_direct_text_json(self, mock_reflector, mock_executor, mock_planner):
        """Test B1: Direct source text provided via JSON payload."""
        payload = {
            "request": "Create an executive proposal for the cloud migration project.",
            "source_text": "Target Cloud Provider: AWS. Database: Aurora PostgreSQL. Target budget: $120,000."
        }
        response = self.client.post("/agent", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("download_url", data)
        sources = data.get("sources", [])
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0]["source_name"], "user_provided_text")
        self.assertEqual(sources[0]["source_type"], "text")
        self.assertIn("Target Cloud Provider: AWS", sources[0]["content"])

    @patch("agent.planner.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.executor.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.reflector.call_llm", side_effect=mock_call_llm_handler)
    def test_scenario_b_direct_text_form(self, mock_reflector, mock_executor, mock_planner):
        """Test B2: Direct source text provided via multipart/form-data."""
        form_data = {
            "request": "Create an SOP for user onboarding in the banking portal.",
            "source_text": "Identity verification requires KYC document upload and SMS OTP confirmation."
        }
        response = self.client.post("/agent", data=form_data)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("download_url", data)
        sources = data.get("sources", [])
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0]["source_type"], "text")
        self.assertIn("Identity verification", sources[0]["content"])

    @patch("agent.planner.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.executor.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.reflector.call_llm", side_effect=mock_call_llm_handler)
    def test_scenario_c_one_txt_file(self, mock_reflector, mock_executor, mock_planner):
        """Test C: Single TXT file upload via multipart/form-data."""
        form_data = {
            "request": "Create an API technical design document for the payment gateway."
        }
        files = {
            "files": ("gateway_spec.txt", b"Endpoints: /v1/payments/charge, /v1/payments/refund.\nAuth: Bearer JWT.\nRate limit: 100 req/min.", "text/plain")
        }
        response = self.client.post("/agent", data=form_data, files=files)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("download_url", data)
        sources = data.get("sources", [])
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0]["source_name"], "gateway_spec.txt")
        self.assertEqual(sources[0]["source_type"], "txt")
        self.assertIn("Bearer JWT", sources[0]["content"])

    @patch("agent.planner.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.executor.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.reflector.call_llm", side_effect=mock_call_llm_handler)
    def test_scenario_d_markdown_file(self, mock_reflector, mock_executor, mock_planner):
        """Test D: Single Markdown file upload via multipart/form-data."""
        form_data = {
            "request": "Create a deployment plan for the microservices cluster."
        }
        files = {
            "files": ("deploy.md", b"# Infrastructure\n- Kubernetes EKS 1.30\n- ArgoCD GitOps\n- Datadog Monitoring", "text/markdown")
        }
        response = self.client.post("/agent", data=form_data, files=files)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("download_url", data)
        sources = data.get("sources", [])
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0]["source_name"], "deploy.md")
        self.assertEqual(sources[0]["source_type"], "md")
        self.assertIn("Kubernetes EKS 1.30", sources[0]["content"])

    @patch("agent.planner.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.executor.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.reflector.call_llm", side_effect=mock_call_llm_handler)
    def test_scenario_e_one_docx_file(self, mock_reflector, mock_executor, mock_planner):
        """Test E: Single DOCX file upload via multipart/form-data."""
        form_data = {
            "request": "Create a project summary report based on our roadmap."
        }
        files = {
            "files": ("roadmap.docx", self.sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        response = self.client.post("/agent", data=form_data, files=files)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("download_url", data)
        sources = data.get("sources", [])
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0]["source_name"], "roadmap.docx")
        self.assertEqual(sources[0]["source_type"], "docx")
        self.assertIn("TechCorp Mobile Banking Roadmap.", sources[0]["content"])

    @patch("agent.planner.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.executor.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.reflector.call_llm", side_effect=mock_call_llm_handler)
    def test_scenario_f_direct_text_plus_file(self, mock_reflector, mock_executor, mock_planner):
        """Test F: Both direct text notes and uploaded document supplied together."""
        form_data = {
            "request": "Draft a comprehensive project launch plan.",
            "source_text": "Meeting Note: Steering committee approved a $20,000 contingency reserve on October 1st."
        }
        files = {
            "files": ("spec.docx", self.sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        response = self.client.post("/agent", data=form_data, files=files)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("download_url", data)
        sources = data.get("sources", [])
        self.assertEqual(len(sources), 2)
        source_names = [s["source_name"] for s in sources]
        self.assertIn("spec.docx", source_names)
        self.assertIn("user_provided_text", source_names)

    def test_validation_empty_request(self):
        """Test validation error for empty or too short request."""
        response = self.client.post("/agent", json={"request": "short"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("minimum 10 characters", response.json()["detail"])

    @patch("agent.planner.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.executor.call_llm", side_effect=mock_call_llm_handler)
    @patch("agent.reflector.call_llm", side_effect=mock_call_llm_handler)
    def test_validation_unsupported_file(self, mock_reflector, mock_executor, mock_planner):
        """Test rejection of unsupported file types."""
        form_data = {"request": "Generate report from binary file."}
        files = {"files": ("virus.exe", b"MZbinaryexec", "application/octet-stream")}
        response = self.client.post("/agent", data=form_data, files=files)
        self.assertEqual(response.status_code, 400)
        self.assertIn("unsupported format", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
