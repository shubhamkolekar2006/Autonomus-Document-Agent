import io
import os
import unittest
from docx import Document
from pypdf import PdfWriter

from agent.document_parser import (
    normalize_text,
    parse_text,
    parse_markdown,
    parse_pdf,
    parse_docx,
    parse_document,
    build_source_context,
    DocumentParsingError,
    UnsupportedFileTypeError,
    FileTooLargeError,
    EmptyDocumentError,
    CorruptedDocumentError,
    MAX_FILE_SIZE_MB
)


class TestDocumentParser(unittest.TestCase):

    def setUp(self):
        # Helper to create valid in-memory PDF
        self.pdf_writer = PdfWriter()
        self.pdf_writer.add_blank_page(width=200, height=200)
        # Note: blank PDF has no text, so we can test EmptyDocumentError on blank PDF,
        # or we can test corrupted/valid PDF extraction.
        self.pdf_stream = io.BytesIO()
        self.pdf_writer.write(self.pdf_stream)
        self.blank_pdf_bytes = self.pdf_stream.getvalue()

        # Helper to create valid in-memory DOCX with text and table
        self.docx_doc = Document()
        self.docx_doc.add_paragraph("Acme Corp Project Specifications.")
        self.docx_doc.add_paragraph("Budget: $75,000. Launch date: November 15, 2026.")
        table = self.docx_doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Phase"
        table.rows[0].cells[1].text = "Timeline"
        table.rows[1].cells[0].text = "Phase 1"
        table.rows[1].cells[1].text = "4 Weeks"
        self.docx_stream = io.BytesIO()
        self.docx_doc.save(self.docx_stream)
        self.valid_docx_bytes = self.docx_stream.getvalue()

    def test_01_txt_extraction(self):
        txt_content = b"  Line 1: Server setup on AWS.\r\n\r\n\r\nLine 2: Python 3.12 and FastAPI.   \n"
        result = parse_text(txt_content)
        self.assertIn("Line 1: Server setup on AWS.", result)
        self.assertIn("Line 2: Python 3.12 and FastAPI.", result)
        self.assertNotIn("\r", result)

    def test_02_markdown_extraction(self):
        md_content = b"# System Architecture\n\n- Component A: Microservices\n- Component B: PostgreSQL\n"
        result = parse_markdown(md_content)
        self.assertIn("# System Architecture", result)
        self.assertIn("Component A: Microservices", result)

    def test_03_pdf_extraction(self):
        # Test valid PDF with readable text or blank pdf error handling
        # For blank PDF:
        with self.assertRaises(EmptyDocumentError):
            parse_pdf(self.blank_pdf_bytes)

    def test_04_docx_extraction(self):
        result = parse_docx(self.valid_docx_bytes)
        self.assertIn("Acme Corp Project Specifications.", result)
        self.assertIn("Budget: $75,000", result)
        self.assertIn("Phase | Timeline", result)

    def test_05_direct_text_normalization(self):
        raw_text = "   Leading spaces \r\n\r\n\r\n Multiple blank lines   \n\n\n Paragraph 2.  \x00 "
        normalized = normalize_text(raw_text)
        self.assertEqual(normalized, "Leading spaces\n\nMultiple blank lines\n\nParagraph 2.")

    def test_06_empty_input(self):
        with self.assertRaises(EmptyDocumentError):
            parse_document("empty.txt", b"")
        with self.assertRaises(EmptyDocumentError):
            parse_document("whitespace.txt", b"   \n\n\t  ")

    def test_07_unsupported_extension(self):
        with self.assertRaises(UnsupportedFileTypeError):
            parse_document("script.exe", b"binary content")
        with self.assertRaises(UnsupportedFileTypeError):
            parse_document("data.csv", b"col1,col2\nval1,val2")

    def test_08_invalid_corrupted_document(self):
        with self.assertRaises(CorruptedDocumentError):
            parse_document("broken.pdf", b"%PDF-1.4 corrupted incomplete stream byte data")
        with self.assertRaises(CorruptedDocumentError):
            parse_document("broken.docx", b"PK\x03\x04 not a valid zip archive file")

    def test_09_file_size_validation(self):
        # 1.5 MB file against 1.0 MB limit
        large_bytes = b"A" * (1024 * 1024 + 500 * 1024)
        with self.assertRaises(FileTooLargeError):
            parse_document("huge.txt", large_bytes, max_file_size_mb=1.0)

    def test_10_multiple_sources(self):
        src1 = parse_document("notes.txt", b"Requirements: Mobile banking with biometric auth.")
        src2 = parse_document("spec.docx", self.valid_docx_bytes)
        
        context_obj = build_source_context(sources=[src1, src2])
        self.assertEqual(len(context_obj["sources"]), 2)
        self.assertIn("SOURCE: notes.txt", context_obj["combined_text"])
        self.assertIn("SOURCE: spec.docx", context_obj["combined_text"])

    def test_11_direct_text_and_uploaded_file_together(self):
        src_file = parse_document("arch.md", b"# Architecture\nMicroservices on AWS ECS.")
        direct_text = "Key stakeholder contact: team@example.com. SLA is 99.9%."
        
        context_obj = build_source_context(sources=[src_file], direct_text=direct_text)
        self.assertEqual(len(context_obj["sources"]), 2)
        self.assertIn("SOURCE: arch.md", context_obj["combined_text"])
        self.assertIn("SOURCE: user_provided_text", context_obj["combined_text"])
        self.assertIn("team@example.com", context_obj["combined_text"])


if __name__ == "__main__":
    unittest.main()
